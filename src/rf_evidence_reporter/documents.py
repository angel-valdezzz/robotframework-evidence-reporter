"""Print formats preserve the same case, milestone and evidence hierarchy as HTML."""

from importlib.resources import files
from io import BytesIO
from xml.sax.saxutils import escape

from .builder import display_date, timezone_label
from .i18n import footer_text, translator


def summary(case, language="en"):
    t = translator(language)
    rows = [
        (t("Suite"), case["suite"]),
        (t("Estatus de ejecución"), case["status"]),
        (t("Inicio"), display_date(case["started_at"], language)),
        (t("Fin"), display_date(case["ended_at"], language)),
        (t("Zona horaria"), timezone_label(case["started_at"], language)),
        (
            t("Duración"),
            f"{case['duration_seconds']:.2f} s"
            if case["duration_seconds"] is not None
            else t("Sin finalizar"),
        ),
    ]
    return [*rows, *case["metadata"].items()]


# Match the light HTML palette; labels remain visible for monochrome printing.
STATUS_PALETTE = {
    "PASS": ("168044", "E5F6EC"),
    "FAIL": ("C52D4B", "FFE8EE"),
    "ERROR": ("C52D4B", "FFE8EE"),
    "WARN": ("A26000", "FFF3C9"),
    "SKIP": ("A26000", "FFF3C9"),
    "INCOMPLETE": ("A26000", "FFF3C9"),
    "INFO": ("1467C2", "E9F2FF"),
}


def event_level(event):
    return event.get("level", "INFO") if event["kind"] == "message" else event.get("status", "INFO")


def status_colors(level):
    return STATUS_PALETTE.get(level, STATUS_PALETTE["INFO"])


def event_text(event):
    return f"{event_level(event)} | {event['title']}"


def printable_image(event):
    # Native print engines consistently support PNG, including transparent WebP sources.
    from PIL import Image

    with Image.open(BytesIO(event["_image_bytes"])) as image:
        image.load()
        content = BytesIO()
        image.save(content, format="PNG")
        content.seek(0)
        return content, image.size


def render_document(case, blocks, target, format_name, *, language="en"):
    if format_name == "pdf":
        render_pdf(case, blocks, target, language=language)
    else:
        render_docx(case, blocks, target, language=language)


def render_pdf(case, blocks, target, *, language="en"):
    t = translator(language)
    from pathlib import Path

    import reportlab
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (
        Image,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    fonts = Path(reportlab.__file__).parent / "fonts"
    pdfmetrics.registerFont(TTFont("EvidenceVera", str(fonts / "Vera.ttf")))
    pdfmetrics.registerFont(TTFont("EvidenceVeraBold", str(fonts / "VeraBd.ttf")))
    pdfmetrics.registerFontFamily(
        "EvidenceVera",
        normal="EvidenceVera",
        bold="EvidenceVeraBold",
        italic="EvidenceVera",
        boldItalic="EvidenceVeraBold",
    )

    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = (
            "EvidenceVeraBold" if "Heading" in style.name or style.name == "Title" else "EvidenceVera"
        )
    styles.add(ParagraphStyle("Evidence", parent=styles["Normal"], fontSize=9, leading=13, spaceAfter=8))
    styles.add(ParagraphStyle("BlockDescription", parent=styles["Evidence"], keepWithNext=True))
    styles["Title"].alignment = TA_LEFT
    styles["Title"].fontSize = 21
    styles["Title"].leading = 26
    styles["Heading1"].fontSize = 15
    styles["Heading1"].leading = 19
    styles["Normal"].leading = 14
    styles.add(
        ParagraphStyle(
            "CardTitle",
            parent=styles["Evidence"],
            fontName="EvidenceVeraBold",
            fontSize=11,
            leading=15,
            spaceAfter=0,
        )
    )
    styles.add(ParagraphStyle("CardBody", parent=styles["Evidence"], spaceAfter=4))
    styles.add(
        ParagraphStyle(
            "Caption",
            parent=styles["Evidence"],
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#52627c"),
            spaceAfter=0,
        )
    )
    story = []

    def paragraph(text, style="Evidence"):
        # Escape markup and allow long unbroken metadata to wrap inside tables.
        return Paragraph(escape(str(text)).replace("\n", "<br/>"), styles[style])

    def panel(title, body, accent, background, *, keep=False):
        # Splittable body supports long messages without overflowing a page.
        title_style = ParagraphStyle(
            "PanelTitle", parent=styles["CardTitle"], textColor=colors.HexColor("#" + accent)
        )
        title_paragraph = Paragraph(escape(str(title)).replace("\n", "<br/>"), title_style)
        card = Table([[title_paragraph], [body]], colWidths=[6.55 * inch], hAlign="LEFT", splitInRow=1)
        card.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + background)),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#" + accent)),
                    ("BOX", (0, 0), (-1, -1), 0.7, colors.HexColor("#" + accent)),
                    ("LINEBEFORE", (0, 0), (0, -1), 3, colors.HexColor("#" + accent)),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 12),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 12),
                    ("TOPPADDING", (0, 0), (-1, -1), 9),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 9),
                ]
            )
        )
        card.keepWithNext = keep
        return card

    story.append(paragraph(case["name"], "Title"))
    accent, background = status_colors(case["status"])
    story.extend(
        [
            panel(
                t(t("Estatus de ejecución")).upper() + "  |  " + case["status"],
                [paragraph(case["message"] or t("Resultado del caso"), "CardBody")],
                accent,
                background,
            ),
            Spacer(1, 12),
        ]
    )
    if case["description"]:
        story.append(paragraph(case["description"]))
    rows = [[paragraph(key), paragraph(value)] for key, value in summary(case, language)]
    table = Table(rows, colWidths=[1.5 * inch, 5.05 * inch], hAlign="LEFT")
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#ede9ff")),
                ("ROWBACKGROUNDS", (1, 0), (-1, -1), [colors.white, colors.HexColor("#f7f9fd")]),
                ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#dce4f1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#d9deeb")),
                ("LEFTPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 7),
            ]
        )
    )
    story.extend([table, Spacer(1, 12)])
    if case.get("attempts"):
        story.append(paragraph(t("Historial de intentos"), "Heading1"))
        for attempt in [*case["attempts"], case]:
            story.append(
                paragraph(
                    f"{display_date(attempt['started_at'])} | {attempt['status']} | {attempt['message']}"
                )
            )
    captures_on_page = 0
    for block in blocks:
        if block.get("id"):
            story.append(PageBreak())
            captures_on_page = 0
        story.append(
            panel(
                block["title"],
                [paragraph(block["description"] or t("Evidencias del caso"), "CardBody")],
                "5744CC",
                "EDE9FF",
                keep=True,
            )
        )
        story.append(Spacer(1, 10))
        for number, event in enumerate(block["events"], 1):
            content = []
            if event.get("description"):
                content.append(paragraph(event["description"], "CardBody"))
            if event.get("_image_bytes"):
                if captures_on_page == 2:
                    story.append(PageBreak())
                    captures_on_page = 0
                captures_on_page += 1
                stream, (width, height) = printable_image(event)
                scale = min((6.55 * inch - 24) / width, 2.65 * inch / height, 1)
                content.append(Image(stream, width=width * scale, height=height * scale, hAlign="LEFT"))
            if event.get("captured_at"):
                content.append(paragraph(display_date(event["captured_at"]), "Caption"))
            if event.get("reason"):
                content.append(paragraph(event["reason"], "CardBody"))
            accent, background = status_colors(event_level(event))
            story.append(
                panel(
                    f"{t('Paso')} {number}  |  {event_text(event)}",
                    content or [paragraph(t("Mensaje registrado"), "CardBody")],
                    accent,
                    background,
                )
            )
            story.append(Spacer(1, 10))

    def footer(canvas, document):
        canvas.saveState()
        from reportlab.lib.utils import ImageReader

        logo = files("rf_evidence_reporter").joinpath("assets/logo.png").read_bytes()
        canvas.drawImage(ImageReader(BytesIO(logo)), 36, A4[1] - 49, width=28, height=28, mask="auto")
        canvas.setFont("EvidenceVeraBold", 10)
        canvas.setFillColor(colors.HexColor("#5744CC"))
        canvas.drawString(74, A4[1] - 39, "Evidence Reporter")
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#52657c"))
        canvas.drawString(36, 23, footer_text(language))
        canvas.drawRightString(A4[0] - 36, 23, f"{t('Página')} {document.page}")
        canvas.restoreState()

    SimpleDocTemplate(
        str(target),
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=70,
        bottomMargin=40,
        title=case["name"],
        author="Evidence Reporter",
    ).build(story, onFirstPage=footer, onLaterPages=footer)


def render_docx(case, blocks, target, *, language="en"):
    t = translator(language)
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.6)
    section.left_margin = section.right_margin = Inches(0.6)
    for name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        if style._element.pPr is not None:
            for border in list(style._element.pPr.findall(qn("w:pBdr"))):
                style._element.pPr.remove(border)
    document.styles["Normal"].font.size = Pt(10)
    document.styles["Normal"].paragraph_format.space_after = Pt(6)
    document.styles["Title"].font.size = Pt(22)

    def shade(cell, color):
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), color)
        cell._tc.get_or_add_tcPr().append(shading)

    def frame(table, color):
        props = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for side in ("top", "left", "bottom", "right"):
            edge = OxmlElement("w:" + side)
            edge.set(qn("w:val"), "single")
            edge.set(qn("w:sz"), "20" if side == "left" else "6")
            edge.set(qn("w:color"), color)
            borders.append(edge)
        props.append(borders)
        margins = OxmlElement("w:tblCellMar")
        for side in ("top", "left", "bottom", "right"):
            edge = OxmlElement("w:" + side)
            edge.set(qn("w:w"), "130" if side in ("top", "bottom") else "180")
            edge.set(qn("w:type"), "dxa")
            margins.append(edge)
        props.append(margins)

    def card(title, accent, background, *, new_page=False):
        table = document.add_table(rows=2, cols=1)
        table.autofit = False
        table.columns[0].width = Inches(7.07)
        frame(table, accent)
        header, body = table.cell(0, 0), table.cell(1, 0)
        # Avoid isolated header rows; Word may still split an oversized body.
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        shade(header, background)
        p = header.paragraphs[0]
        p.paragraph_format.page_break_before = new_page
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(title)
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor.from_string(accent)
        return body

    def body_paragraph(cell, text):
        p = (
            cell.paragraphs[0]
            if not cell.paragraphs[0].text and not cell.paragraphs[0].runs
            else cell.add_paragraph()
        )
        p.add_run(str(text))
        return p

    brand = document.add_paragraph()
    brand.add_run().add_picture(
        BytesIO(files("rf_evidence_reporter").joinpath("assets/logo.png").read_bytes()),
        width=Inches(0.32),
    )
    brand.add_run("  Evidence Reporter")
    brand.paragraph_format.keep_with_next = True
    brand.runs[-1].bold = True
    brand.runs[-1].font.color.rgb = RGBColor.from_string("5744CC")
    document.add_paragraph(case["name"], "Title")
    if case["description"]:
        document.add_paragraph(case["description"])
    accent, background = status_colors(case["status"])
    result = card(t(t("Estatus de ejecución")).upper() + "  |  " + case["status"], accent, background)
    body_paragraph(result, case["message"] or t("Resultado del caso"))
    document.add_paragraph().paragraph_format.space_after = Pt(2)
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Inches(1.5)
    table.columns[1].width = Inches(5.57)
    for key, value in summary(case, language):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = str(key), str(value)
        cells[0].width, cells[1].width = Inches(1.5), Inches(5.5)
        shade(cells[0], "EDE9FF")
        if len(table.rows) % 2 == 0:
            shade(cells[1], "F7F9FD")
        if key == t("Estatus de ejecución"):
            shade(cells[1], background)
            for run in cells[1].paragraphs[0].runs:
                run.font.color.rgb = RGBColor.from_string(accent)
                run.bold = True
        for run in cells[0].paragraphs[0].runs:
            run.bold = True
    document.add_paragraph()
    if case.get("attempts"):
        document.add_heading(t("Historial de intentos"), 1)
        for attempt in [*case["attempts"], case]:
            document.add_paragraph(
                f"{display_date(attempt['started_at'])} | {attempt['status']} | {attempt['message']}"
            )
    captures_on_page = 0
    for block in blocks:
        if block.get("id"):
            # Separate adjacent tables so Word keeps each milestone as its own card.
            document.add_paragraph().paragraph_format.space_after = Pt(2)
        heading = card(block["title"], "5744CC", "EDE9FF", new_page=bool(block.get("id")))
        if block.get("id"):
            captures_on_page = 0
        body_paragraph(
            heading, block["description"] or t("Evidencias del caso")
        ).paragraph_format.keep_with_next = True
        gap = document.add_paragraph()
        gap.paragraph_format.space_after = Pt(2)
        gap.paragraph_format.keep_with_next = True
        for number, event in enumerate(block["events"], 1):
            new_page = False
            if event.get("_image_bytes"):
                if captures_on_page == 2:
                    new_page = True
                    captures_on_page = 0
                captures_on_page += 1
            accent, background = status_colors(event_level(event))
            body = card(
                f"{t('Paso')} {number}  |  {event_text(event)}", accent, background, new_page=new_page
            )
            if event.get("description"):
                body_paragraph(body, event["description"])
            if event.get("_image_bytes"):
                stream, (width, height) = printable_image(event)
                scale = min(6.8 / width, 2.65 / height, 1 / 96)
                p = body_paragraph(body, "")
                p.add_run().add_picture(stream, width=Inches(width * scale), height=Inches(height * scale))
            if event.get("captured_at"):
                p = body_paragraph(body, display_date(event["captured_at"]))
                for run in p.runs:
                    run.font.size = Pt(8)
                    run.font.color.rgb = RGBColor.from_string("52627C")
            if event.get("reason"):
                body_paragraph(body, event["reason"])
            for p in body.paragraphs[:-1]:
                p.paragraph_format.keep_with_next = True
            if number < len(block["events"]):
                document.add_paragraph().paragraph_format.space_after = Pt(2)
    footer = section.footer.paragraphs[0]
    footer.add_run(footer_text(language) + " | " + t("Página") + " ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.core_properties.title = case["name"]
    document.core_properties.author = "Evidence Reporter"
    document.save(target)
